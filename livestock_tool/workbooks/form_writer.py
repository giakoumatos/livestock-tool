"""Writes the data collection workbook: a "Start here" sheet, one Dataset
sheet per dataset (workbooks/form.py lays it out), and the hidden lists of
the dropdowns. With no datasets it writes the empty template; with the
records of a project (as form.read_sheet returns them) it writes them in
(tools/convert_workbook.py)."""
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import form
from .form import (CLASSES, COL_DATA, COL_HELP, COL_KEY, COL_LABEL, COL_NEED, COL_NO, COL_SOURCE,
                   COL_UNIT, FORM, INPUT_COLUMNS, INPUTS, MANURE, N_CLASSES, N_INPUTS, N_PRODUCTS,
                   PRODUCTS, SINGLE, key)
from .schema import CH, EU, OPTIONAL, RECOMMENDED, REQUIRED, SHEETS

VERSION = '3.0'
FONT = 'Arial'
RED, BAND, PART, INK, MUTED = 'C8102E', '2F3B48', 'E6EAF0', '1F2937', '6B7280'
INPUT, IDLE, LINE = 'FFF6CC', 'F3F4F6', 'D0D5DD'
NEED_COLOURS = ((REQUIRED, 'B42318'), (RECOMMENDED, '067647'), (OPTIONAL, MUTED), (EU, '175CD3'),
                (CH, '175CD3'))
CONDITIONAL = 'B54708'
WIDTHS = {COL_KEY: 28, COL_NO: 5, COL_LABEL: 48, COL_UNIT: 15, COL_NEED: 18, COL_SOURCE: 36,
          COL_HELP: 64}
DATA_WIDTH = 18
MAX_LINES = 6           # a filled cell's longer text shows in the formula bar
_SIDE = Side(style='thin', color=LINE)
BOX = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)
WRAP = Alignment(wrap_text=True, vertical='top')
_LAST = COL_HELP

START = 'Start here'
HOW = (
    'Fill one Dataset sheet per dataset. For a second dataset, right-click the "Dataset 1" tab, '
    'choose Move or Copy, tick "Create a copy" and rename the copy "Dataset 2".',
    'Enter data only in the yellow cells. Where a cell has a dropdown arrow, choose from the list. '
    'Grey cells are not used for that row.',
    'Fill every row marked Required, then the parts that apply to your dataset (table below). Leave '
    'the other rows empty: the tool then uses the cited default, or does not need the value.',
    'Give the source of each value in its Source cell (for products and inputs: in their comment '
    'or source row). The comments of the dataset are written from them.',
    'Amounts are for the whole herd and one year unless the unit says otherwise; the tool divides '
    'them by the reference product (the first product column of section 6).',
    'Run: python run_tool.py projects/<project>/<workbook>.xlsx. A filled example: '
    'projects/cashmere/livestock_input_cashmere.xlsx.',
)
PARTS_TO_FILL = (
    ('Every dataset', 'Sections 1-3; section 4 "Herd and feed"; section 5 shares (not Switzerland); '
     'section 6 "Product" and "Price"; section 7; section 9.'),
    ('Classes without feed intake', 'Section 4 "Tier 2 gross energy" (cattle: or the simplified '
     'intake equation under "Cattle"). With no Tier 2 data either (not cattle): section 2 '
     '"Tier 1 region".'),
    ('Cattle and buffalo', 'Section 4 "Cattle".'),
    ('Pigs and poultry', 'Section 4 "Pigs and poultry".'),
    ('Milk, fibre or animals sold', 'Section 6 "Milk", "Fibre", "Live weight" (dairy herds also '
     '"Dairy live weight").'),
    ('Liquid manure or a digester', 'Section 5 "Storage details".'),
    ('EU-27 country', 'The rows marked "EU only" (sections 4 and 5).'),
    ('Switzerland', 'The rows marked "CH only" (sections 4 and 5); the manure shares of section 5 '
     'are not used.'),
    ('Water or irrigation rows in section 7', 'Section 8.'),
)
METHODS = (
    ('Switzerland', 'SALCAanimal V1.0 (SALCA, Agroscope), as documented in its appendix.'),
    ('EU-27', 'IPCC 2019 Refinement for greenhouse gases and N flows; EMEP/EEA 2023 for NH3, NOx, '
     'NMVOC and particulate matter.'),
    ('Any other country', 'IPCC 2019 Refinement, within the scope of the WFLDB v3.5 guidelines.'),
)
LEGEND = (
    (INPUT, 'Enter data here.'),
    (IDLE, 'Not used for this row.'),
)
NEED_LEGEND = (
    ('Required', 'B42318', 'Needed for every dataset (or, as the tag says, every dataset except '
     'Switzerland, or unless the alternative is given).'),
    ('Recommended', '067647', 'A cited default is used when empty, but the value moves the results.'),
    ('Optional', MUTED, 'Empty = the cited default, or the value is not needed.'),
    ('Tier 2, Cattle, Dairy, ...', CONDITIONAL, 'Only for that case: classes without feed intake, '
     'cattle, dairy, a product type, a digester, water.'),
    ('EU only, CH only', '175CD3', 'Only for datasets in the EU-27, or in Switzerland.'),
)
# The inputs that move the results most (sensitivity on the cashmere farmed
# herd, +10 % on each input; Tier 2 figures from a 36 kg cashmere doe).
SIGNIFICANT = (
    ('Feed intake (or Tier 2 data)', '4: dry matter intake (or Tier 2 gross energy)',
     'Sets the gross energy, hence enteric methane, volatile solids and N intake: +10 % intake '
     'gives +10 % methane, N2O and NH3.'),
    ('Ym (methane conversion)', '4: Ym (cattle: Ym class)', '+10 % Ym gives +9.9 % methane, about '
     '+9 % of the greenhouse gases.'),
    ('Crude protein of the ration', '4: crude protein', '+10 % crude protein gives +10 % N '
     'excreted, so +10 % of every N emission.'),
    ('Wet or dry climate', '2: wet or dry climate', 'Dry to wet raised N2O by 39 % (EF4 0.005 to '
     '0.014; N leaching 0 to 0.24 of pasture N).'),
    ('Manure management and housing', '5: shares; 4: days in buildings', '+10 % pasture share '
     'gave -7.4 % N2O; liquid systems raise manure methane strongly (MCF up to 80 %).'),
    ('Digestibility of the feed', '4: DE', 'Tier 2: +10 % DE gives -12.7 % gross energy; with the '
     'intake given it only sets the volatile solids.'),
    ('Herd size and products', '4: animals present; 6: amounts and prices', 'All results are '
     'divided by the reference product; prices set the allocation.'),
)


def write_workbook(path, datasets=()):
    wb = Workbook()
    start = wb.active
    start.title = START
    _start_sheet(start)
    lists = _Lists(wb.create_sheet('lists'))
    for n, records in enumerate(datasets or [None], 1):
        ws = wb.create_sheet(f'Dataset {n}', n)
        _dataset_sheet(ws, lists)
        if records:
            _write_records(ws, records)
    lists.ws.sheet_state = 'hidden'
    wb.active = 0
    wb.save(path)


# --- styling helpers --------------------------------------------------------

def _font(size=10, bold=False, italic=False, color=INK):
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color)


def _fill(color):
    return PatternFill('solid', fgColor=color)


def _band(ws, r, color, height=None):
    for c in range(COL_NO, _LAST + 1):
        ws.cell(row=r, column=c).fill = _fill(color)
    if height:
        ws.row_dimensions[r].height = height


def _text(ws, r, c, value, **font):
    cell = ws.cell(row=r, column=c, value=value)
    cell.font = _font(**font)
    cell.alignment = WRAP
    return cell


def _merge(ws, r, c1, c2):
    if c2 > c1:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)


def _lines(text, width):
    """Rows of text a wrapped cell of this width needs."""
    if not text:
        return 1
    return sum(max(1, -(-len(part) // max(1, int(width * 1.1)))) for part in str(text).split('\n'))


def _banner(ws, title, subtitle):
    _band(ws, 1, RED, 30)
    _band(ws, 2, RED, 20)
    _text(ws, 1, COL_NO, title, size=16, bold=True, color='FFFFFF').alignment = \
        Alignment(vertical='center')
    _text(ws, 2, COL_NO, subtitle, size=10, color='FFFFFF').alignment = Alignment(vertical='top')
    _merge(ws, 1, COL_NO, _LAST)
    _merge(ws, 2, COL_NO, _LAST)


def _widths(ws):
    for c in range(1, _LAST + 1):
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS.get(c, DATA_WIDTH)
    ws.column_dimensions[get_column_letter(COL_KEY)].hidden = True
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# --- Start here -------------------------------------------------------------

def _start_sheet(ws):
    _widths(ws)
    ws.column_dimensions[get_column_letter(COL_LABEL)].width = 34
    _banner(ws, 'Livestock LCI data collection', f'Data collection template for the ecoinvent '
            f'livestock tool, version {VERSION}')
    r = 4

    def heading(title):
        nonlocal r
        r += 1
        _band(ws, r, BAND, 22)
        _text(ws, r, COL_NO, title, size=12, bold=True, color='FFFFFF').alignment = \
            Alignment(vertical='center')
        _merge(ws, r, COL_NO, _LAST)
        r += 1

    def line(label, text, label_color=INK, fill=None):
        nonlocal r
        cell = _text(ws, r, COL_LABEL, label, bold=True, color=label_color)
        if fill:
            cell.fill = _fill(fill)
            cell.border = BOX
        _text(ws, r, COL_UNIT, text)
        _merge(ws, r, COL_UNIT, _LAST)
        ws.row_dimensions[r].height = 14 * max(_lines(text, 150), _lines(label, 30))
        r += 1

    heading('How to fill this workbook')
    for n, text in enumerate(HOW, 1):
        _text(ws, r, COL_NO, n, bold=True).alignment = Alignment(horizontal='center',
                                                                  vertical='top')
        _text(ws, r, COL_LABEL, text)
        _merge(ws, r, COL_LABEL, _LAST)
        ws.row_dimensions[r].height = 14 * _lines(text, 200)
        r += 1
    heading('Which parts to fill')
    for label, text in PARTS_TO_FILL:
        line(label, text)
    heading('How to read the Dataset sheet')
    for color, text in LEGEND:
        line('', text, fill=color)
    for label, color, text in NEED_LEGEND:
        line(label, text, label_color=color)
    heading('Methods by geography (set by the country or region in section 1)')
    for label, text in METHODS:
        line(label, text)
    heading('The inputs that matter most')
    for label, where, why in SIGNIFICANT:
        line(label, f'Section {where}. {why}')
    ws.freeze_panes = 'A4'


# --- Dataset sheet ------------------------------------------------------------

class _Lists:
    """The hidden sheet of dropdown values, one column per list."""

    def __init__(self, ws):
        self.ws, self.columns, self.validations = ws, {}, {}

    def validation(self, sheet_ws, choices):
        choices = tuple(str(c) for c in choices)
        if choices not in self.columns:
            n = len(self.columns) + 1
            for i, v in enumerate(choices, 1):
                self.ws.cell(row=i, column=n, value=v)
            self.columns[choices] = get_column_letter(n)
        k = (sheet_ws.title, choices)
        if k not in self.validations:
            letter = self.columns[choices]
            dv = DataValidation(type='list', allow_blank=True,
                                formula1=f'=lists!${letter}$1:${letter}${len(choices)}')
            dv.error = 'Choose a value from the list.'
            dv.errorStyle = 'warning'
            sheet_ws.add_data_validation(dv)
            self.validations[k] = dv
        return self.validations[k]


def _dataset_sheet(ws, lists):
    _widths(ws)
    _banner(ws, 'Livestock LCI data collection: dataset', 'Enter data in the yellow cells. Fill '
            'the rows marked Required and the parts that apply to this dataset (see "Start here").')
    ws.cell(row=1, column=COL_KEY, value=form.MARKER)
    r = 4
    for section in FORM:
        r += 1
        _band(ws, r, BAND, 24)
        _text(ws, r, COL_NO, section.number, size=12, bold=True, color='FFFFFF').alignment = \
            Alignment(horizontal='center', vertical='center')
        _text(ws, r, COL_LABEL, section.title, size=12, bold=True, color='FFFFFF').alignment = \
            Alignment(vertical='center')
        if section.note:
            _text(ws, r, COL_UNIT, section.note, size=9, italic=True, color='FFFFFF').alignment = \
                Alignment(wrap_text=True, vertical='center')
            _merge(ws, r, COL_UNIT, _LAST)
            ws.row_dimensions[r].height = max(24, 13 * _lines(section.note, 180))
        r += 1
        for part in section.parts:
            r = _part(ws, lists, r, part)
    ws.freeze_panes = ws.cell(row=4, column=COL_DATA)


def _captions(kind):
    if kind in (CLASSES, MANURE):
        return [f'Class {i + 1}' for i in range(N_CLASSES)]
    if kind == PRODUCTS:
        return ['Reference product'] + [f'By-product {i}' for i in range(1, N_PRODUCTS)]
    return ['Data']


def _part(ws, lists, r, part):
    _band(ws, r, PART, 18)
    if part.kind == INPUTS:
        for field, heading, c1, c2 in INPUT_COLUMNS:
            _text(ws, r, c1, heading, bold=True)
            _merge(ws, r, c1, c2)
        _text(ws, r, COL_HELP, 'How to fill', bold=True)
        ws.row_dimensions[r].height = 28
        r += 1
        r = _note(ws, r, _inputs_note())
        columns = {f: c for f, _, c, _ in INPUT_COLUMNS}
        spans = {f: c2 for f, _, _, c2 in INPUT_COLUMNS}
        fields = {c.name: c for c in SHEETS['inputs']}
        for n in range(1, N_INPUTS + 1):
            ws.cell(row=r, column=COL_KEY, value=key(INPUTS, n))
            _text(ws, r, COL_NO, n, size=8, color=MUTED).alignment = Alignment(horizontal='center')
            for f, c in columns.items():
                _input(ws, r, c, spans[f])
                if fields[f].choices:
                    lists.validation(ws, fields[f].choices).add(ws.cell(row=r, column=c))
            r += 1
        return r
    _text(ws, r, COL_LABEL, part.title or 'Question', bold=True)
    _text(ws, r, COL_UNIT, 'Unit', bold=True)
    _text(ws, r, COL_NEED, 'Needed', bold=True)
    captions = _captions(part.kind)
    for i, caption in enumerate(captions):
        _text(ws, r, COL_DATA + i, caption, bold=True)
    if part.kind == SINGLE:
        _merge(ws, r, COL_DATA, COL_SOURCE - 1)
    _text(ws, r, COL_SOURCE, 'Source', bold=True)
    _text(ws, r, COL_HELP, 'How to fill', bold=True)
    if part.kind == PRODUCTS:
        ws.row_dimensions[r].height = 28
    r += 1
    if part.note:
        r = _note(ws, r, part.note)
    for row in part.rows:
        col = form.column(part.kind, row.field)
        source = form.source_column(part.kind, row)
        ws.cell(row=r, column=COL_KEY, value=key(part.kind, row.field))
        _text(ws, r, COL_LABEL, row.label)
        _text(ws, r, COL_UNIT, '' if col.unit in ('-', 'text') else col.unit, size=9, color=MUTED)
        if part.kind != MANURE:      # the part's note says what the systems need
            _text(ws, r, COL_NEED, form.need_tag(col.need), size=9,
                  bold=col.need.startswith(REQUIRED), color=_need_colour(col.need))
        if part.kind == SINGLE:
            cells = [_input(ws, r, COL_DATA, COL_SOURCE - 1)]
        else:
            n = N_PRODUCTS if part.kind == PRODUCTS else N_CLASSES
            cells = [_input(ws, r, COL_DATA + i) for i in range(n)]
            for i in range(n, N_CLASSES):
                ws.cell(row=r, column=COL_DATA + i).fill = _fill(IDLE)
        if col.choices:
            dv = lists.validation(ws, col.choices)
            for cell in cells:
                dv.add(cell)
        if isinstance(col.example, str) and col.unit == 'yyyy-mm-dd':
            for cell in cells:
                cell.number_format = 'yyyy-mm-dd'
        if source is not None:
            _input(ws, r, COL_SOURCE)
        else:
            ws.cell(row=r, column=COL_SOURCE).fill = _fill(IDLE)
        if part.kind == MANURE:
            help_text = ''
        else:
            help_text = col.text + (f' Example: {col.example}.' if col.example and not col.choices
                                    else '')
            if source is not None:
                help_text += f' Source: {source.text[0].lower() + source.text[1:]}'
        _text(ws, r, COL_HELP, help_text, size=9, color=MUTED)
        long_text = col.unit == 'text' and part.kind == SINGLE
        ws.row_dimensions[r].height = max(
            16, 12 * _lines(help_text, WIDTHS[COL_HELP]), 12 * _lines(row.label, WIDTHS[COL_LABEL]),
            48 if long_text else 0)
        r += 1
    return r


def _note(ws, r, text):
    _text(ws, r, COL_LABEL, text, size=9, italic=True, color=MUTED)
    _merge(ws, r, COL_LABEL, _LAST)
    ws.row_dimensions[r].height = max(15, 12 * _lines(text, 230))
    return r + 1


def _inputs_note():
    return ('Type: FromTechnosphere = bought or supplied inputs (feed, energy, materials, '
            'services); FromEnvironment = resources (water, land occupation); ToEnvironment = '
            'emissions the models do not compute (e.g. veterinary products). Basis: herd_year = '
            'for the herd and the year (the tool divides by the reference product); per_kg = '
            'already per kg of the reference product. Compartment and subcompartment: environment '
            'rows only. Water use: animals (drinking and cleaning) or irrigation.')


def _input(ws, r, c1, c2=None):
    for c in range(c1, (c2 or c1) + 1):
        cell = ws.cell(row=r, column=c)
        cell.fill = _fill(INPUT)
        cell.border = BOX
        cell.font = _font()
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    _merge(ws, r, c1, c2 or c1)
    return ws.cell(row=r, column=c1)


def _need_colour(need):
    for start, colour in NEED_COLOURS:
        if need.startswith(start):
            return colour
    return CONDITIONAL


# --- filling a project's data -------------------------------------------------

def _write_records(ws, records):
    rows = {ws.cell(row=r, column=COL_KEY).value: r for r in range(1, ws.max_row + 1)}
    spans = {(m.min_row, m.min_col): sum(WIDTHS.get(c, DATA_WIDTH)
                                         for c in range(m.min_col, m.max_col + 1))
             for m in ws.merged_cells.ranges}

    def put(k, c, value):
        if value is None or value == '':
            return
        r = rows[k]
        cell = ws.cell(row=r, column=c)
        cell.value = value
        if isinstance(value, (date, datetime)):
            cell.number_format = 'yyyy-mm-dd'
        elif isinstance(value, str):       # a long text: the row grows to show it
            width = spans.get((r, c), WIDTHS.get(c, DATA_WIDTH))
            ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 15,
                                              13 * min(MAX_LINES, _lines(value, width)))

    a = records['activity']
    herd = records['herd']
    if len(herd) > N_CLASSES:
        raise ValueError(f'{a.get("activityName")}: {len(herd)} animal classes, the sheet has '
                         f'{N_CLASSES}')
    products = sorted(records['products'], key=lambda p: p.get('group') != 'ReferenceProduct')
    if len(products) > N_PRODUCTS:
        raise ValueError(f'{a.get("activityName")}: {len(products)} products, the sheet has '
                         f'{N_PRODUCTS}')
    if len(records['inputs']) > N_INPUTS:
        raise ValueError(f'{a.get("activityName")}: {len(records["inputs"])} inputs, the sheet has '
                         f'{N_INPUTS}')
    for k, (kind, row) in form.INDEX.items():
        if kind == SINGLE:
            put(k, COL_DATA, a.get(row.field))
            if row.source:
                put(k, COL_SOURCE, a.get(row.source))
        elif kind == CLASSES:
            for i, h in enumerate(herd):
                put(k, COL_DATA + i, h.get(row.field))
            if row.source:
                put(k, COL_SOURCE, _joined([(h.get('class_name'), h.get(row.source))
                                            for h in herd]))
        elif kind == MANURE:
            found = []
            for i, h in enumerate(herd):
                mine = [m for m in records['manure'] if m['class_name'] == h.get('class_name')] or \
                    [m for m in records['manure'] if m['class_name'] == 'all']
                shares = [m for m in mine if m['system'] == row.field]
                if shares:
                    put(k, COL_DATA + i, sum(m['share'] for m in shares))
                    found += [(h.get('class_name'), m.get('source')) for m in shares]
            put(k, COL_SOURCE, _joined(found))
        elif kind == PRODUCTS:
            for i, p in enumerate(products):
                put(k, COL_DATA + i, _product_value(records, p, row.field))
    for n, item in enumerate(records['inputs'], 1):
        for f, _, c, _ in INPUT_COLUMNS:
            put(key(INPUTS, n), c, item.get(f))


def _product_value(records, product, field):
    props = {p['propertyName']: p for p in records['properties']
             if p['name'] == product.get('name')}
    if field == form.PRICE:
        return props.get('price', {}).get('amount')
    if field == form.PRICE_UNIT:
        return props.get('price', {}).get('unitName')
    if field == form.PRICE_COMMENT:
        return props.get('price', {}).get('comment')
    if field == form.PROPERTIES_SOURCE:
        return _joined([(None, p.get('comment')) for name, p in props.items() if name != 'price'])
    if field in form.PROPERTY_UNITS:
        return props.get(field, {}).get('amount')
    return product.get(field)


def _joined(pairs):
    """One source text from those of several classes: the shared text, or
    each class's text after its name when they differ."""
    texts = list(dict.fromkeys(t for _, t in pairs if t))
    if len(texts) <= 1:
        return texts[0] if texts else None
    by_text = {}
    for name, t in pairs:
        if t:
            by_text.setdefault(t, []).append(name)
    return '; '.join(f'{", ".join(n for n in names if n)}: {t}' if any(names) else t
                     for t, names in by_text.items())
