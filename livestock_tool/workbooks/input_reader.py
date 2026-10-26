"""Reads the filled livestock data collection workbook into datasets (the
crop tool's ExcelInputReader step): every Dataset sheet (workbooks/form.py)
gives one dataset."""
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook

from ..models.base import fmt
from ..models.model_enums import CATEGORIES, SYSTEMS
from ..tables.idf2022 import fpcm
from ..tables.ipcc2019 import CLIMATE_TREE_WHERE, MOISTURE_WHERE, climate_zone, \
    moisture_from_zone
from .form import MARKER, read_sheet

_NUMBER = {'head', 'dmi_kg_day', 'dmi_kg_year_class', 'de_pct', 'cp_pct', 'ash_frac', 'ue_frac',
           'ym_pct', 'nex_kg_year', 'housing_days', 'weight_kg', 'weight_start_kg',
           'weight_end_kg', 'mature_weight_kg', 'weight_gain_kg_day', 'milk_kg_day',
           'milk_fat_pct', 'wool_kg_year', 'pregnant_share', 'offspring_per_birth',
           'silage_share_of_max', 'share', 'amount_per_year', 'amount', 'liquid_retention_months',
           'farms_sampled', 'share_of_national_production', 'production_volume', 'fat_pct',
           'protein_pct', 'weaning_gain_kg', 'ndf_pct', 'forage_share', 'grain_share',
           'water_release_share', 'irrigation_efficiency', 'straw_kg_year', 'yard_share', 'ge_mj_per_kg_dm',
           'pasture_days_spring', 'pasture_days_summer', 'pasture_days_autumn',
           'pasture_days_winter', 'pasture_days', 'pasture_hours', 'yard_days', 'yard_hours',
           'grazing_n_ratio', 'liquid_n_share', 'lw_purchased_kg_year_class',
           'lw_sold_kg_year_class', 'lw_gain_kg_year_class', 'milk_kg_year_class',
           'eggs_kg_year_class', 'liquid_crust_share', 'liquid_pit_share', 'liquid_digester_share',
           'solid_compost_share', 'solid_deep_bedding_share', 'mat_c', 'map_mm', 'pet_mm',
           'elevation_m', 'frost_days', 'warmest_month_c', 'cover_ch4_reduction',
           'frac_gas_digester', 'production_days', 'litter_size', 'farrowings_per_year',
           'piglet_birth_kg', 'piglet_weaning_kg', 'sow_weight_change_kg', 'egg_g_day',
           'nemf_mj_kg', 'lactation_days', 'scoured_yield', 'grease_share',
           'body_ebw', 'body_protein', 'body_fat', 'body_ash', 'egg_shell_share'}
_YES_NO = {'irrigated_pasture', 'pasture_in_dataset', 'tied_housing', 'crust_thick_dry'}
SHARE_TOLERANCE = 0.02
# Units the ecoinvent flows do not use: converted to kg (any row) or m3 (rows
# from or to the environment, i.e. water), with a note in the row's comment.
_MASS_UNITS = {'mg': 1e-6, 'g': 1e-3, 'ug': 1e-9, 'µg': 1e-9, 't': 1e3, 'tonne': 1e3}
_VOLUME_UNITS = {'l': 1e-3, 'ml': 1e-6}
# Open points that need someone to act (a choice or data the dataset lacks)
# start with this; all others record how a gap was handled.
ACTION = 'ACTION: '


@dataclass
class Dataset:
    activity: dict
    herd: list = field(default_factory=list)
    manure: list = field(default_factory=list)
    products: list = field(default_factory=list)
    properties: list = field(default_factory=list)
    inputs: list = field(default_factory=list)
    open_points: list = field(default_factory=list)

    @property
    def name(self):
        return self.activity['activityName']

    def reference_product(self):
        ref = [p for p in self.products if p['group'] == 'ReferenceProduct']
        if len(ref) != 1:
            raise ValueError(f'{self.name}: needs exactly one ReferenceProduct, has {len(ref)}')
        return ref[0]

    @staticmethod
    def amount(product):
        """A product's yearly amount as the dataset counts it: FPCM for
        cattle and buffalo milk (WFLDB v3.5 section 2.3; IDF 2022)."""
        return product.get('fpcm_kg') or product['amount_per_year']

    def manure_rows(self, class_name):
        """The "manure" rows of a class, or those of "all"."""
        return [m for m in self.manure if m['class_name'] == class_name] or \
            [m for m in self.manure if m['class_name'] == 'all']

    def shares_for(self, class_name):
        rows = self.manure_rows(class_name)
        if not rows:
            raise ValueError(f'{self.name}: no manure shares for "{class_name}"')
        shares = {}
        for m in rows:
            if m['system'] not in SYSTEMS:
                raise ValueError(f'{self.name}: unknown manure system "{m["system"]}"')
            shares[m['system']] = shares.get(m['system'], 0.0) + m['share']
        total = sum(shares.values())
        if abs(total - 1) > SHARE_TOLERANCE:
            raise ValueError(f'{self.name}, {class_name}: manure shares add up to {fmt(total)}')
        if abs(total - 1) > 1e-9:
            shares = {k: v / total for k, v in shares.items()}
            point = (f'{class_name}: the manure shares add up to {fmt(total * 100, 3)} %, so they '
                     f'were scaled to 100 %.')
            if point not in self.open_points:
                self.open_points.append(point)
        return shares


def _cell(name, value):
    if isinstance(value, str):
        value = value.strip()
        if value == '':
            return None
    if value is None:
        return None
    if name in _YES_NO:
        return str(value).lower() in ('yes', 'y', 'true', '1')
    if name in _NUMBER:
        return float(value)
    if isinstance(value, datetime):
        return value.date()
    return value


def _typed(record):
    return {k: _cell(k, v) for k, v in record.items()}


def _dataset_sheets(wb, path):
    sheets = [ws for ws in wb.worksheets
              if next(ws.iter_rows(max_row=1, max_col=1, values_only=True), (None,))[0] == MARKER]
    if not sheets:
        if 'activity' in wb.sheetnames:
            raise ValueError(f'{path}: this is the old table layout; convert it with '
                             f'python tools/convert_workbook.py "{path}" <new.xlsx>')
        raise ValueError(f'{path}: no Dataset sheet (copy "Dataset 1" of the template)')
    return sheets


def _ecoinvent_unit(item):
    """Converts an input row in mg, g, t (or L of water) to kg (or m3); the
    note for the open points, or None."""
    unit = str(item.get('unitName') or '').strip()
    env = item.get('group') in ('FromEnvironment', 'ToEnvironment')
    if unit.lower() in _MASS_UNITS:
        target, factor = 'kg', _MASS_UNITS[unit.lower()]
    elif env and unit.lower() in _VOLUME_UNITS:
        target, factor = 'm3', _VOLUME_UNITS[unit.lower()]
    else:
        return None
    old = item['amount']
    item['amount'], item['unitName'] = old * factor, target
    note = f'Converted from {fmt(old)} {unit} to {fmt(old * factor)} {target}.'
    item['comment'] = ' '.join(c for c in (item.get('comment'), note) if c)
    return f'{item["name"]}: {note[0].lower() + note[1:-1]}, the unit of ecoinvent.'


def _feeding_days(text):
    if not text:
        return None
    days = {}
    for part in str(text).split(';'):
        if part.strip():
            k, v = part.split(':')
            days[k.strip()] = float(v)
    return days


def _climate(a):
    """Climate zone and wet/dry class from the climate data when they are
    left empty (IPCC 2019 Vol. 4, Ch. 3, Figure 3A.5.2; Ch. 11, Table 11.1,
    fn 4)."""
    source = a.get('climate_data_source') or 'the climate data given'
    if a.get('climate_zone') is None and a.get('mat_c') is not None:
        a['climate_zone'] = climate_zone(a['mat_c'], a.get('map_mm'), a.get('pet_mm'),
                                         a.get('elevation_m'), a.get('frost_days'),
                                         a.get('warmest_month_c'))
        data = [f'mean annual temperature {fmt(a["mat_c"])} deg C']
        if a.get('map_mm') is not None:
            data.append(f'precipitation {fmt(a["map_mm"])} mm')
        if a.get('pet_mm') is not None:
            data.append(f'PET {fmt(a["pet_mm"])} mm')
        a['climate_zone_source'] = f'{CLIMATE_TREE_WHERE}, from {source} ({", ".join(data)})'
    if a.get('climate_zone') is None:
        raise ValueError(f'{a["activityName"]}: give the IPCC climate zone or the climate data '
                         f'(section 2)')
    if a.get('moisture') is None:
        a['moisture'] = moisture_from_zone(a['climate_zone'], a.get('map_mm'))
        a['moisture_source'] = f'{MOISTURE_WHERE}, from the climate zone {a["climate_zone"]}'


def read_datasets(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    datasets = []
    for ws in _dataset_sheets(wb, path):
        raw = read_sheet(ws)
        a = _typed(raw['activity'])
        if not a.get('activityName'):
            raise ValueError(f'sheet "{ws.title}": section 1 has no dataset name')
        if any(d.name == a['activityName'] for d in datasets):
            raise ValueError(f'sheet "{ws.title}": a second dataset named "{a["activityName"]}"')
        if a.get('pasture_in_dataset') is None:
            raise ValueError(f'{a["activityName"]}: answer "Does this dataset include the grazing '
                             f'land?" (section 3, yes or no); WFLDB v3.5 counts the excreta on '
                             f'pasture in the pasture datasets (section 3.12.2)')
        _climate(a)
        for k in ('startDate', 'endDate'):
            if isinstance(a[k], (date, datetime)):
                a[k] = a[k].isoformat()[:10]
        name = a['activityName']
        mine = {s: [_typed(r) for r in raw[s]] for s in raw if s != 'activity'}
        for h in mine['herd']:
            if h['category'] not in CATEGORIES:
                raise ValueError(f'{name}: unknown category "{h["category"]}"')
            if h['dmi_kg_day'] is None and h['dmi_kg_year_class'] is not None:
                h['dmi_kg_day'] = h['dmi_kg_year_class'] / 365 / h['head']
            h['feeding_days'] = _feeding_days(h['feeding_days'])
        bovine = any(h['category'] in ('dairy_cattle', 'buffalo') for h in mine['herd'])
        for p in mine['products']:
            if (bovine and p['product_type'] == 'milk' and p['fat_pct'] is not None
                    and p['protein_pct'] is not None):
                p['fpcm_kg'] = fpcm(p['amount_per_year'], p['fat_pct'], p['protein_pct'])
        points = [n for n in map(_ecoinvent_unit, mine['inputs']) if n]
        datasets.append(Dataset(a, mine['herd'], mine['manure'], mine['products'],
                                mine['properties'], mine['inputs'], points))
    return datasets
