"""Writes the ecoinvent data provision workbook (sheets meta, facts,
exchanges, properties, parameters), the format the agri-lca-pipeline reads
from its "input datasets" folder. The two tools meet only in this file."""
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

META_COLUMNS = (
    'activityName', 'geography', 'specialActivityType', 'startDate', 'endDate', 'technologyLevel',
    'synonyms', 'tags', 'includedActivitiesStart', 'includedActivitiesEnd', 'generalComment',
    'geographyComment', 'allocationComment', 'technologyComment', 'timePeriodComment', 'ISIC',
    'dataEntryBy_isActiveAuthor', 'dataEntryBy_personName',
    'dataGeneratorAndPublication_accessRestrictedTo', 'dataGeneratorAndPublication_dataPublishedIn',
    'dataGeneratorAndPublication_pageNumbers', 'dataGeneratorAndPublication_personName',
    'dataGeneratorAndPublication_publishedSourceId', 'modellingAndValidation_extrapolations',
    'modellingAndValidation_percent', 'modellingAndValidation_samplingProcedure',
    'fileAttributes_majorRelease', 'fileAttributes_minorRelease', 'fileAttributes_majorRevision',
    'fileAttributes_minorRevision')
FACTS_COLUMNS = (
    'activityName', 'data_source', 'data_source_reference', 'farms_sampled', 'extrapolations',
    'region_detail', 'share_of_national_production', 'pv_years', 'yield_source', 'price_source')
_UNCERTAINTY = (
    'uncertainty type', 'uncertainty comment', 'variance', 'minValue', 'mostLikelyValue',
    'maxValue', 'reliability', 'completeness', 'temporalCorrelation', 'geographicalCorrelation',
    'furtherTechnologyCorrelation')
EXCHANGE_COLUMNS = (
    'activityName', 'geography', 'group', 'name', 'compartment', 'subcompartment',
    'activityLink_activityName', 'activityLink_geography', 'unitName', 'amount', 'variableName',
    'mathematicalRelation', 'comment', 'pageNumbers', 'sourceId', *_UNCERTAINTY,
    'PV_amount', 'PV_variableName', 'PV_mathematicalRelation', 'PV_comment', 'PV_sourceId',
    *('PV_' + c for c in _UNCERTAINTY))
PROPERTY_COLUMNS = (
    'activityName', 'geography', 'group', 'name', 'compartment', 'subcompartment',
    'activityLink_activityName', 'activityLink_geography', 'propertyName', 'unitName', 'amount',
    'variableName', 'mathematicalRelation', 'comment', *_UNCERTAINTY)
PARAMETER_COLUMNS = (
    'activityName', 'geography', 'name', 'unitName', 'amount', 'variableName',
    'mathematicalRelation', 'comment', *_UNCERTAINTY)


@dataclass
class ProvisionDataset:
    meta: dict
    facts: dict = field(default_factory=dict)
    exchanges: list = field(default_factory=list)
    properties: list = field(default_factory=list)
    parameters: list = field(default_factory=list)


def write_provision(path, datasets):
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, 'meta', META_COLUMNS, [d.meta for d in datasets])
    _sheet(wb, 'facts', FACTS_COLUMNS, [d.facts for d in datasets if d.facts])
    _sheet(wb, 'exchanges', EXCHANGE_COLUMNS, [r for d in datasets for r in d.exchanges])
    _sheet(wb, 'properties', PROPERTY_COLUMNS, [r for d in datasets for r in d.properties])
    _sheet(wb, 'parameters', PARAMETER_COLUMNS, [r for d in datasets for r in d.parameters])
    wb.save(path)


def _sheet(wb, title, columns, rows):
    ws = wb.create_sheet(title)
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        unknown = set(row) - set(columns)
        if unknown:
            raise KeyError(f'{title}: unknown columns {sorted(unknown)}')
        ws.append([row.get(c) for c in columns])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=False, vertical='top')
