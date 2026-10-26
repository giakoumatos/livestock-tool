"""Converts a data collection workbook in the old table layout (sheets
activity, herd, manure, products, properties, inputs, one row per record)
to the questionnaire layout: one Dataset sheet per dataset.

    python tools/convert_workbook.py old.xlsx [new.xlsx]

Without new.xlsx the old file is replaced (it is read in full first).
"""
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from livestock_tool.workbooks.form import check_coverage  # noqa: E402
from livestock_tool.workbooks.form_writer import write_workbook  # noqa: E402

SHEETS = ('activity', 'herd', 'manure', 'products', 'properties', 'inputs')


def _records(wb, sheet):
    if sheet not in wb.sheetnames:
        return []
    rows = wb[sheet].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else None for h in next(rows)]
    out = []
    for raw in rows:
        if raw and isinstance(raw[0], str) and raw[0].startswith('[explanation]'):
            continue
        record = {h: (v.strip() or None) if isinstance(v, str) else v
                  for h, v in zip(header, raw) if h}
        if any(v is not None for v in record.values()):
            out.append(record)
    return out


def convert(source, target):
    wb = load_workbook(source, data_only=True, read_only=True)
    tables = {s: _records(wb, s) for s in SHEETS}
    wb.close()
    datasets = []
    for a in tables['activity']:
        name = a['activityName']
        datasets.append({'activity': a, **{s: [r for r in tables[s] if r.get('activityName') == name]
                                           for s in SHEETS if s != 'activity'}})
    if not datasets:
        raise ValueError(f'{source}: no rows on "activity"; is it the old table layout?')
    write_workbook(target, datasets)
    return datasets


def main():
    if len(sys.argv) not in (2, 3):
        sys.exit(__doc__)
    check_coverage()
    source = Path(sys.argv[1])
    target = Path(sys.argv[2]) if len(sys.argv) == 3 else source
    datasets = convert(source, target)
    print(f'{target}: {len(datasets)} Dataset sheets '
          f'({", ".join(d["activity"]["activityName"] for d in datasets)})')


if __name__ == '__main__':
    main()
